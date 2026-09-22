from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(__file__).with_name("Nexora_Part_2_Literature_Review_Shortened.docx")


SECTIONS = [
    (
        "2.1 Geopolitical risk and tourism demand",
        [
            "Geopolitical risk covers threats and realised events involving war, terrorism and interstate tension. Caldara and Iacoviello's (2022) news-based index distinguishes these dimensions. For tourism, such risk can alter perceived safety, travel advisories, air routes, operating costs and destination choice, including when conflict occurs outside Malaysia. Global conflicts may also affect household confidence and energy prices before changing observed travel demand.",
            "Evidence is generally negative but heterogeneous. Lee et al. (2021) found lower long-run tourism demand across 16 countries, yet no causal relationship between GPR and Malaysian tourist numbers; Malaysian tourism receipts showed a different pattern. Papagianni et al. (2024) also reported persistent negative responses in most of 14 emerging economies, with stronger historical effects in China, Indonesia and Thailand. The size, timing and direction of the response therefore cannot be assumed uniform across destinations or source markets.",
            "Malaysia combines neighbouring land markets with long-haul markets dependent on air connectivity. National totals may conceal differences in income, trip purpose and exposure to particular events. This study consequently models monthly arrivals from 20 source markets and includes global GPR, Malaysia-specific GPR and country interactions to allow the controlled association to vary across markets. This design can distinguish a broad international shock from a response concentrated in selected origins.",
            "Because GPR also moves with economic conditions, energy prices, post-pandemic recovery and omitted factors, prediction does not establish causation. The model uses GPR as a forecasting signal, while Ridge coefficients and SHAP values are interpreted as model-based associations only.",
        ],
    ),
    (
        "2.2 Oil Prices, Transport Costs and Tourism Demand",
        [
            "Transport is part of a trip's total price, and oil affects both aviation and road costs. Higher transport expenditure can also reduce the budget available at the destination. Becken and Lennox (2012) linked sustained oil-price increases to lower tourism exports through transport, income and exchange-rate channels. Across 23 European countries, Boto-Garcia (2025) estimated an average arrival elasticity near -0.2 and persistent effects from supply-driven oil shocks. Although context-specific, this supports energy prices as short-term predictors.",
            "Oil nevertheless imperfectly represents traveller costs because airlines hedge fuel, vary surcharges and adjust capacity, while land travellers may face regulated retail prices. Seetaram (2010) warns that oil, distance and published airfares introduce measurement error. Tan and Soon (2023) likewise found that Malaysian travel-cost responses differed by region rather than showing a uniform demand reduction.",
            "In 2024, 41.0% of Malaysian tourists arrived by air and 52.4% by land (Tourism Malaysia, 2025). The model therefore combines Brent crude as global transport-cost pressure with Malaysian retail fuel prices as a road-travel indicator. This captures two possible transmission channels without assuming that crude-price changes reach every market equally or immediately. The distinction is especially relevant to Malaysia's mix of regional land arrivals and long-haul aviation demand.",
        ],
    ),
    (
        "2.3 Exchange Rates and Malaysia's Destination Affordability",
        [
            "Exchange rates determine how much visitors can buy in Malaysia with origin-country income. Ringgit depreciation should improve affordability, whereas appreciation reduces price competitiveness. Volatility can also create uncertainty about the final trip cost, delaying bookings, shortening stays or encouraging destination substitution. Tourism-demand reviews identify income, destination prices and bilateral exchange rates as core variables (Song & Li, 2008), while disaggregated evidence shows that level and volatility effects differ across origins (Imamboccus et al., 2024).",
            "Malaysian evidence rejects a simple weaker-ringgit-equals-more-tourists rule. Karimi et al. (2019) found asymmetric effects, with both appreciations and depreciations reducing arrivals in some specifications, which they linked to price rigidity and substitution. Across 21 origins, Tan and Soon (2023) found depreciation increased ASEAN demand but reduced demand from China, other Asian markets and Western countries. Income, airfares, competing destinations and incomplete price pass-through may further shape these responses.",
            "Bilateral rates are therefore preferable to an aggregate index: one ringgit movement changes affordability differently for visitors paid in Singapore dollars, yuan or US dollars. The pooled model preserves this variation while sharing information across markets. Exchange rates are treated as market-specific affordability indicators rather than isolated causal determinants. This avoids masking opposing origin-market responses.",
        ],
    ),
    (
        "2.4 Malaysian Tourism Resilience, Economic Conditions and SDG Alignment",
        [
            "Malaysia recorded 25,016,698 tourist arrivals in 2024, 24.2% above 2023 but 4.2% below 2019 (Tourism Malaysia, 2025). In 2025, tourism generated RM323.0 billion in value added, or 15.9% of GDP, and supported 3.7 million jobs (Department of Statistics Malaysia [DOSM], 2026). These figures show both the recovery and the sector's economic scale. Forecast errors therefore affect staffing, transport, marketing and the timing of public support.",
            "Recovery has not removed vulnerability. COVID-19 exposed labour precarity and complex supply-chain weaknesses in Penang (Hampton et al., 2023), while dependence on a few source markets increases exposure to health, economic and geopolitical shocks. Country-level forecasts reveal whether weakness is broad or concentrated, enabling earlier adjustments to marketing, connectivity and workforce plans.",
            "This decision-support role supports the National Tourism Policy 2020-2030 emphasis on competitiveness, sustainability, inclusiveness, smart tourism and responsible tourism (Ministry of Tourism, Arts and Culture Malaysia [MOTAC], 2020). It also relates to SDG target 8.9, covering sustainable tourism, jobs and local culture, and target 12.b, which calls for tools that monitor tourism's sustainability impacts (United Nations Department of Economic and Social Affairs [UN DESA], n.d.).",
            "Forecasting itself does not prove environmental or social sustainability. Its contribution is operational: timely market-level information can reduce avoidable over- or under-capacity and support planning for volatility. The dashboard should show forecast coverage and limitations, while decisions combine predictions with environmental, community and labour indicators rather than equating higher arrivals with sustainable progress. This is important because tourism performance affects both national output and a large tourism-related workforce.",
        ],
    ),
    (
        "2.5 Machine-Learning Approaches and Research Gap",
        [
            "Tourism demand is seasonal, persistent and shock-sensitive. Traditional methods include naive and seasonal-naive forecasts, exponential smoothing, autoregressive models and econometric equations. Machine learning can capture nonlinear interactions among lags, seasonality, prices and external indicators, but no method dominates every destination or horizon (Song & Li, 2008). Forecast accuracy also varies with the origin, destination, sample, frequency, variables and horizon (Peng et al., 2014), so models require task-specific out-of-sample testing.",
            "Pooling lets short market series share information through one global forecasting function. Global models can generalise across related series and are not inherently more restrictive than separate local models (Montero-Manso & Hyndman, 2021). However, tourism accuracy suffers if market heterogeneity is ignored; fixed, spatial or temporal effects can improve performance (Long et al., 2019). This study therefore adds country identifiers, market features and GPR-by-country interactions rather than assuming identical responses.",
            "LightGBM captures nonlinear thresholds and mixed-feature interactions (Ke et al., 2017) and performed effectively with correlated series and exogenous variables in the M5 competition (Makridakis et al., 2022). However, retail success does not guarantee tourism superiority. Ridge regression provides a transparent regularised comparison, while naive-last and seasonal-naive forecasts remain essential benchmarks. This comparison tests whether added complexity improves on simple historical rules. SHAP assigns feature contributions to predictions but explains model behaviour rather than causal effects (Lundberg & Lee, 2017).",
            "Rolling-origin evaluation reproduces repeated real-time decisions and measures each lead time without random-split leakage (Tashman, 2000). Separate horizon results are important because accuracy may decline as the forecast extends. MAE and RMSE retain arrival scale; MAPE can become unstable for small markets with near-zero values (Hyndman & Koehler, 2006).",
            "The applied gap is that Malaysian research mainly estimates aggregate or long-run relationships, while forecasts rarely integrate monthly source-market arrivals, bilateral exchange rates, global and Malaysian GPR, energy prices, publication lags and post-COVID stress tests. This project fills that gap with a pooled direct model for 20 markets and one- to four-month horizons, tested against transparent benchmarks using only information available at each forecast origin. Direct horizon-specific models also avoid recursively feeding earlier prediction errors into later horizons. External indicators are not presumed to help every horizon; their value is evaluated empirically.",
        ],
    ),
]


REFERENCES = [
    "Becken, S., & Lennox, J. (2012). Implications of a long-term increase in oil prices for tourism. Tourism Management, 33(1), 133-142. https://doi.org/10.1016/j.tourman.2011.02.012",
    "Boto-Garcia, D. (2025). The elasticity of tourist arrivals to oil prices. Tourism Economics. Advance online publication. https://doi.org/10.1177/13548166251364305",
    "Caldara, D., & Iacoviello, M. (2022). Measuring geopolitical risk. American Economic Review, 112(4), 1194-1225. https://doi.org/10.1257/aer.20191823",
    "Department of Statistics Malaysia. (2026, September 15). Tourism satellite account 2025. https://www.dosm.gov.my/portal-main/release-content/tourism-satellite-account-2025",
    "Hampton, M. P., Jeyacheya, J., & Nair, V. (2023). Post-COVID tourism revealed: Evidence from Malaysia. Annals of Tourism Research, 103, 103671. https://doi.org/10.1016/j.annals.2023.103671",
    "Hyndman, R. J., & Koehler, A. B. (2006). Another look at measures of forecast accuracy. International Journal of Forecasting, 22(4), 679-688. https://doi.org/10.1016/j.ijforecast.2006.03.001",
    "Imamboccus, R., Seetanah, B., Nunkoo, R., & Jaffur, Z. K. (2024). The impact of exchange rate and exchange rate volatility on tourism demand using disaggregated data. International Journal of Tourism Research, 26(2), e2640. https://doi.org/10.1002/jtr.2640",
    "Karimi, M. S., Khan, A. A., & Karamelikli, H. (2019). Asymmetric effects of real exchange rate on inbound tourist arrivals in Malaysia: An analysis of price rigidity. International Journal of Tourism Research, 21(2), 156-164. https://doi.org/10.1002/jtr.2249",
    "Ke, G., Meng, Q., Finley, T., Wang, T., Chen, W., Ma, W., Ye, Q., & Liu, T.-Y. (2017). LightGBM: A highly efficient gradient boosting decision tree. In Advances in Neural Information Processing Systems 30 (pp. 3146-3154). Curran Associates.",
    "Lee, C.-C., Olasehinde-Williams, G., & Akadiri, S. S. (2021). Geopolitical risk and tourism: Evidence from dynamic heterogeneous panel models. International Journal of Tourism Research, 23(1), 26-38. https://doi.org/10.1002/jtr.2389",
    "Long, W., Liu, C., & Song, H. (2019). Pooling in tourism demand forecasting. Journal of Travel Research, 58(7), 1161-1174. https://doi.org/10.1177/0047287518800390",
    "Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. In Advances in Neural Information Processing Systems 30 (pp. 4765-4774). Curran Associates.",
    "Makridakis, S., Spiliotis, E., & Assimakopoulos, V. (2022). M5 accuracy competition: Results, findings, and conclusions. International Journal of Forecasting, 38(4), 1346-1364. https://doi.org/10.1016/j.ijforecast.2021.11.013",
    "Ministry of Tourism, Arts and Culture Malaysia. (2020). National Tourism Policy 2020-2030. https://www.motac.gov.my/en/category/download/dasar-pelancongan-negara-dpn-2020-2030/",
    "Montero-Manso, P., & Hyndman, R. J. (2021). Principles and algorithms for forecasting groups of time series: Locality and globality. International Journal of Forecasting, 37(4), 1632-1653. https://doi.org/10.1016/j.ijforecast.2021.03.004",
    "Papagianni, E., Evgenidis, A., Tsagkanos, A., & Megalooikonomou, V. (2024). Tourism demand in the face of geopolitical risk: Insights from a cross-country analysis. Journal of Travel Research, 63(8), 2094-2119. https://doi.org/10.1177/00472875231206539",
    "Peng, B., Song, H., & Crouch, G. I. (2014). A meta-analysis of international tourism demand forecasting and implications for practice. Tourism Management, 45, 181-193. https://doi.org/10.1016/j.tourman.2014.04.005",
    "Seetaram, N. (2010). Computing airfare elasticities or opening Pandora's box. Research in Transportation Economics, 26(1), 27-36. https://doi.org/10.1016/j.retrec.2009.10.005",
    "Song, H., & Li, G. (2008). Tourism demand modelling and forecasting: A review of recent research. Tourism Management, 29(2), 203-220. https://doi.org/10.1016/j.tourman.2007.07.016",
    "Tan, C.-Y., & Soon, S.-V. (2023). Tourism demand for Malaysia: Further evidence from panel approaches. Asia Pacific Management Review, 28(4), 459-469. https://doi.org/10.1016/j.apmrv.2022.12.006",
    "Tashman, L. J. (2000). Out-of-sample tests of forecasting accuracy: An analysis and review. International Journal of Forecasting, 16(4), 437-450. https://doi.org/10.1016/S0169-2070(00)00065-0",
    "Tourism Malaysia. (2025). Visitor performance to Malaysia January to December 2024 [Infographic]. https://data.tourism.gov.my/frontend/pdf/Infographic_visitor_performance_to_malaysia_jan-dec_2024.pdf",
    "United Nations Department of Economic and Social Affairs. (n.d.). Sustainable tourism. United Nations Sustainable Development. https://sdgs.un.org/topics/sustainable-tourism",
]


def set_cell_margins(cell, top=100, start=100, bottom=100, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.5)
sec.bottom_margin = Cm(1.5)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)

for style_name in ("Heading 1", "Heading 2"):
    style = styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.color.rgb = None
    style.font.bold = True
    style.paragraph_format.keep_with_next = True
    style.paragraph_format.space_before = Pt(6)
    style.paragraph_format.space_after = Pt(2)

styles["Heading 1"].font.size = Pt(12)
styles["Heading 2"].font.size = Pt(12)

if "Reference" not in styles:
    ref_style = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
else:
    ref_style = styles["Reference"]
ref_style.font.name = "Times New Roman"
ref_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
ref_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
ref_style.font.size = Pt(12)
ref_style.paragraph_format.line_spacing = 1.5
ref_style.paragraph_format.left_indent = Cm(1.27)
ref_style.paragraph_format.first_line_indent = Cm(-1.27)
ref_style.paragraph_format.space_after = Pt(0)
ref_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

p = doc.add_paragraph(style="Heading 1")
p.paragraph_format.space_before = Pt(0)
set_font(p.add_run("2.0 LITERATURE REVIEW"), bold=True)

for heading, paragraphs in SECTIONS:
    p = doc.add_paragraph(style="Heading 2")
    set_font(p.add_run(heading), bold=True)
    for text in paragraphs:
        p = doc.add_paragraph(style="Normal")
        set_font(p.add_run(text.strip()))

doc.add_page_break()
p = doc.add_paragraph(style="Heading 1")
p.paragraph_format.space_before = Pt(0)
set_font(p.add_run("REFERENCES"), bold=True)

for ref in REFERENCES:
    p = doc.add_paragraph(style="Reference")
    set_font(p.add_run(ref), size=12)

doc.core_properties.title = "Part 2 Literature Review for the Nexora DOSM Datathon 2026 Report"
doc.core_properties.subject = "Literature review on geopolitical risk, transport costs, exchange rates, tourism resilience, sustainability, and machine-learning forecasting"
doc.core_properties.author = "Nexora"
doc.save(OUT)
print(OUT)
